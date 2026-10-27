
from __future__ import annotations

import json
import re
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any

import pandas as pd
import pdfplumber
from docx import Document

SUPPORTED_FORMATS = {
    ".pdf", ".xlsx", ".csv", ".docx", ".txt", ".eml", ".json"
}

COLUMN_ALIASES = {
    "name": {
        "vendor", "vendor_name", "supplier", "supplier_name", "name",
        "bidder", "bidder_name", "company", "company_name",
        "firm_name", "manufacturer", "manufacturer_name",
        "quotation_from", "quoted_by", "agency_name"
    },
    "price": {
        "price", "unit_price", "unit_price_inr", "price_inr",
        "cost", "amount", "price_rs", "total_price", "quoted_price",
        "rate", "quoted_amount", "total_amount", "bid_amount",
        "financial_bid", "grand_total", "total_cost",
        "basic_price", "ex_showroom_price", "on_road_price",
        "vehicle_cost", "item_rate", "evaluated_price"
    },
    "ram_gb": {
        "ram", "ram_gb", "memory", "memory_gb"
    },
    "ssd_gb": {
        "ssd", "ssd_gb", "storage", "storage_gb"
    },
    "delivery_days": {
        "delivery", "delivery_days", "lead_time", "lead_time_days",
        "delivery_lead_time_days", "delivery_period_days",
        "completion_period_days", "supply_period_days"
    },
    "warranty_years": {
        "warranty", "warranty_years", "warranty_yrs",
        "warranty_period", "warranty_period_years"
    },
}


def _number(value: Any) -> float | int | None:
    """Convert an explicitly supplied value to a number."""
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        return value if pd.notna(value) else None

    text = str(value).strip().lower().replace(",", "")
    if not text:
        return None

    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None

    number = float(match.group())
    return int(number) if number.is_integer() else number


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized_keys = {
        re.sub(r"[^a-z0-9]+", "_", str(key).strip().lower()).strip("_"): value
        for key, value in row.items()
    }

    result: dict[str, Any] = {}

    for field, aliases in COLUMN_ALIASES.items():
        value = None

        for key, candidate in normalized_keys.items():
            if key in aliases and candidate is not None and str(candidate).strip():
                value = candidate
                break

        if field == "name":
            result[field] = str(value).strip() if value is not None else None
        else:
            result[field] = _number(value)

    # Convert explicitly labelled TB storage values to GB.
    for field in ("ram_gb", "ssd_gb"):
        raw_value = next(
            (
                str(value).strip().lower()
                for key, value in normalized_keys.items()
                if key in COLUMN_ALIASES[field]
                and value is not None
                and str(value).strip()
            ),
            "",
        )

        if result[field] is not None and re.search(r"\btb\b", raw_value):
            result[field] *= 1024

    # Support an explicit warranty_months column.
    warranty_months = _number(normalized_keys.get("warranty_months"))
    if result["warranty_years"] is None and warranty_months is not None:
        result["warranty_years"] = warranty_months / 12
    elif result["warranty_years"] is not None:
        warranty_raw = str(normalized_keys.get("warranty", "")).lower()
        if "month" in warranty_raw:
            result["warranty_years"] /= 12

    # Keep every other column too (material, load_capacity_kg, ...), so
    # category-specific rules can be evaluated later.
    known_keys = {key for aliases in COLUMN_ALIASES.values() for key in aliases}
    known_keys.add("warranty_months")
    unit_suffixes = (
        "_kg", "_gb", "_days", "_years", "_mm", "_cm", "_inr", "_pct",
        "_percent", "_count", "_qty",
    )

    for key, value in normalized_keys.items():
        if not key or key in known_keys or key in result:
            continue
        if value is None:
            continue

        if isinstance(value, (list, dict)):
            result[key] = value
            continue

        text_value = str(value).strip()
        if not text_value:
            continue

        is_plain_number = re.fullmatch(r"-?\d+(?:\.\d+)?", text_value.replace(",", ""))
        if is_plain_number or key.endswith(unit_suffixes):
            number = _number(value)
            result[key] = number if number is not None else text_value
        else:
            result[key] = text_value

    result["missing_fields"] = [
        field for field in COLUMN_ALIASES if result.get(field) is None
    ]
    result["complete"] = not result["missing_fields"]
    return result


    

def _extract_text_fields(text: str) -> dict[str, Any]:
    """Extract quotation details from common document layouts."""

    text = text.lstrip("\ufeff")
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    result: dict[str, Any] = {
        "name": None,
        "price": None,
        "ram_gb": None,
        "ssd_gb": None,
        "delivery_days": None,
        "warranty_years": None,
    }

    # 1. Prefer an explicitly named vendor.
    vendor_pattern = re.compile(
        r"^\s*(?:vendor|supplier|bidder|company|firm)"
        r"(?:\s+name)?\s*[:\-]\s*(.+?)\s*$",
        re.IGNORECASE,
    )

    # If vendor is not named, use the product/model as a display fallback.
    product_pattern = re.compile(
        r"^\s*(?:proposed\s+model|vehicle\s+model|product\s+name|"
        r"item\s+description|product|item|proposed\s+vehicle)"
        r"\s*[:\-]\s*(.+?)\s*$",
        re.IGNORECASE,
    )

    for line in lines:
        match = vendor_pattern.search(line)
        if match:
            result["name"] = match.group(1).strip()
            break

    if not result["name"]:
        for line in lines:
            match = product_pattern.search(line)
            if match:
                result["name"] = match.group(1).strip()
                break

    # 2. Extract price from labelled lines.
    # Prefer unit price over total quotation value.
    price_patterns = [
        r"\bunit\s+price(?:\s*\([^)]*\))?\s*[:\-]\s*"
        r"(?:INR|Rs\.?|₹)?\s*₹?\s*([\d,]+(?:\.\d+)?)",

        r"\b(?:quoted|basic|vehicle|ex[\s-]?showroom|on[\s-]?road)"
        r"\s+price\s*[:\-]\s*(?:INR|Rs\.?|₹)?\s*₹?\s*"
        r"([\d,]+(?:\.\d+)?)",

        r"\bprice(?:\s*\([^)]*\))?\s*[:\-]\s*"
        r"(?:INR|Rs\.?|₹)?\s*₹?\s*([\d,]+(?:\.\d+)?)",

        r"\b(?:total\s+quoted\s+value|grand\s+total|"
        r"total\s+price|total\s+amount|quoted\s+amount)"
        r"\s*[:\-]\s*(?:INR|Rs\.?|₹)?\s*₹?\s*"
        r"([\d,]+(?:\.\d+)?)",
    ]

    for pattern in price_patterns:
        for line in lines:
            match = re.search(pattern, line, re.IGNORECASE)
            if match:
                result["price"] = _number(match.group(1))
                break
        if result["price"] is not None:
            break

    # 3. Extract computer specifications when present.
    for line in lines:
        if result["ram_gb"] is None:
            match = re.search(
                r"\b(?:ram|memory)\s*[:\-]?\s*"
                r"(\d+(?:\.\d+)?)\s*(gb)?\b",
                line,
                re.IGNORECASE,
            )
            if match:
                result["ram_gb"] = _number(match.group(1))

        if result["ssd_gb"] is None:
            match = re.search(
                r"\b(?:ssd|storage)\s*[:\-]?\s*"
                r"(\d+(?:\.\d+)?)\s*(gb|tb)?\b",
                line,
                re.IGNORECASE,
            )
            if match:
                result["ssd_gb"] = _number(match.group(1))
                if (match.group(2) or "").lower() == "tb":
                    result["ssd_gb"] *= 1024

        if result["delivery_days"] is None:
            match = re.search(
                r"\b(?:delivery(?:\s+period|\s+lead\s+time)?|"
                r"lead\s*time|supply\s+period|completion\s+period)"
                r"\s*[:\-]?\s*(?:within\s+)?(\d+)\s*days?\b",
                line,
                re.IGNORECASE,
            )
            if match:
                result["delivery_days"] = _number(match.group(1))

        if result["warranty_years"] is None:
            match = re.search(
                r"\bwarranty(?:\s+period)?\s*[:\-]?\s*"
                r"(\d+(?:\.\d+)?)\s*(years?|yrs?|months?)",
                line,
                re.IGNORECASE,
            )
            if match:
                result["warranty_years"] = _number(match.group(1))
                if match.group(2).lower().startswith("month"):
                    result["warranty_years"] /= 12

    result["missing_fields"] = [
        field for field in COLUMN_ALIASES
        if result.get(field) is None
    ]
    result["complete"] = not result["missing_fields"]

    return result

    for field, pattern in patterns.items():
        match = re.search(
            pattern,
            text,
            flags=re.IGNORECASE | re.MULTILINE,
        )

        if match is None:
            result[field] = None
            continue

        result[field] = (
            match.group(1).strip()
            if field == "name"
            else _number(match.group(1))
        )

        if field == "ssd_gb" and (match.group(2) or "").lower() == "tb":
            result[field] *= 1024

        if (
            field == "warranty_years"
            and (match.group(2) or "").lower().startswith("month")
        ):
            result[field] /= 12

    result["missing_fields"] = [
        field for field in COLUMN_ALIASES if result.get(field) is None
    ]
    result["complete"] = not result["missing_fields"]

    return result




def _read_text_file(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    if suffix == ".docx":
        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        return "\n".join(parts)

    if suffix == ".pdf":
        pages = []

        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")

                for table in page.extract_tables():
                    for row in table:
                        pages.append(" | ".join(cell or "" for cell in row))

        return "\n".join(pages)

    if suffix == ".eml":
        with path.open("rb") as file:
            message = BytesParser(policy=policy.default).parse(file)

        parts = []

        if message.is_multipart():
            for part in message.walk():
                if part.get_content_type() == "text/plain":
                    content = part.get_content()
                    if isinstance(content, str):
                        parts.append(content)
        elif message.get_content_type() == "text/plain":
            content = message.get_content()
            if isinstance(content, str):
                parts.append(content)

        return "\n".join(parts)

    raise ValueError(f"Text extraction is not supported for {suffix}")


def parse_quotation(file_path: str | Path) -> dict[str, Any]:
    """
    Parse a quotation into normalized vendor records.

    Supported: PDF, XLSX, CSV, DOCX, TXT, EML and JSON.
    Missing fields are reported rather than invented.
    """
    path = Path(file_path)

    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"Quotation file not found: {path}")

    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_FORMATS:
        raise ValueError(
            f"Unsupported format: {suffix}. "
            f"Supported formats: {', '.join(sorted(SUPPORTED_FORMATS))}"
        )

    vendors: list[dict[str, Any]] = []
    warnings: list[str] = []
    raw_text = ""

    if suffix in {".csv", ".xlsx"}:
        frame = (
            pd.read_csv(path)
            if suffix == ".csv"
            else pd.read_excel(path)
        )

        records = frame.astype(object).where(pd.notna(frame), None).to_dict(
            orient="records"
        )
        vendors = [_normalize_row(row) for row in records]

    elif suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))

        if isinstance(data, dict):
            rows = data.get("vendors", [data])
        elif isinstance(data, list):
            rows = data
        else:
            raise ValueError(
                "JSON must contain an object or a list of vendor objects."
            )

        if not isinstance(rows, list) or not all(
            isinstance(row, dict) for row in rows
        ):
            raise ValueError("JSON vendors must be a list of objects.")

        vendors = [_normalize_row(row) for row in rows]

    else:
        raw_text = _read_text_file(path)

        if not raw_text.strip():
            warnings.append(
                "No extractable text found. This may be a scanned PDF "
                "or an email without a plain-text body."
            )
        else:
            record = _extract_text_fields(raw_text)
            vendors = [record]

            if not record["complete"]:
                warnings.append(
                    "Some fields could not be extracted confidently; "
                    "manual review is required."
                )

    if not vendors:
        warnings.append("No vendor records were extracted.")

    return {
        "source_file": path.name,
        "format": suffix.lstrip("."),
        "vendors": vendors,
        "raw_text": raw_text,
        "warnings": warnings,
    }

def test_csv_supports_vehicle_quotation(tmp_path):
    file_path = tmp_path / "vehicle.csv"
    pd.DataFrame([{
        "Supplier Name": "ABC Motors",
        "Product": "Electric Bus",
        "Unit Price INR": 2500000,
        "Delivery Lead Time Days": 30,
    }]).to_csv(file_path, index=False)

    result = parse_quotation(file_path)
    vendor = result["vendors"][0]

    assert vendor["name"] == "ABC Motors"
    assert vendor["price"] == 2500000


def test_csv_supports_furniture_quotation(tmp_path):
    file_path = tmp_path / "chairs.csv"
    pd.DataFrame([{
        "Supplier": "Office Furniture Co",
        "Item Description": "Ergonomic Chair",
        "Amount": 4500,
        "Quantity": 20,
    }]).to_csv(file_path, index=False)

    result = parse_quotation(file_path)
    vendor = result["vendors"][0]

    assert vendor["name"] == "Office Furniture Co"
    assert vendor["price"] == 4500


def test_generic_quotation_preserves_missing_fields(tmp_path):
    file_path = tmp_path / "generic.csv"
    pd.DataFrame([{
        "Vendor": "Reliable Supplies",
        "Product": "Office Desk",
        "Price": 8000,
    }]).to_csv(file_path, index=False)

    result = parse_quotation(file_path)
    vendor = result["vendors"][0]

    assert vendor["name"] == "Reliable Supplies"
    assert vendor["price"] == 8000
    assert vendor["complete"] is False
    assert vendor["missing_fields"]
